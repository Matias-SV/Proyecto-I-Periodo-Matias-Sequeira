### Importamos las herramientas necesarias de Flask para manejar rutas, plantillas, sesiones y redirecciones. ###
from flask import Flask, render_template, request, redirect, url_for, session
import openpyxl
import os

### Inicializamos la aplicación Flask y definimos la clave secreta para el manejo de sesiones. ###
app = Flask(__name__)
app.secret_key = 'clave_secreta_777'

### Indicamos las credenciales de acceso al sistema y nombre del archivo Excel donde se almacenan los contactos. ###
USUARIO = "admin"
PASSWORD = "1234"
EXCEL = "contactos.xlsx"


### Creamos una función para leer todos los contactos desde el archivo Excel, retornando una lista de diccionarios. ###
def leer_contactos():
    if not os.path.exists(EXCEL):
        return []
    wb = openpyxl.load_workbook(EXCEL)
    ws = wb.active
    lista = []
    ### Iteramos desde la segunda fila para omitir el encabezado, asignando un id basado en el índice de la fila. ###
    for i, fila in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if fila and fila[0]:
            lista.append({
                "id": i,
                "nombre": str(fila[0]) if fila[0] else "",
                "apellido": str(fila[1]) if fila[1] else "",
                "telefono": str(fila[2]) if fila[2] else "",
                "correo": str(fila[3]) if fila[3] else "",
                "direccion": str(fila[4]) if fila[4] else "",
                "categoria": str(fila[5]) if fila[5] else "Otro",
                "favorito": str(fila[6]) if fila[6] else "no"
            })
    return lista


### Creamos una función para guardar la lista completa de contactos en el archivo Excel, sobreescribiendo el contenido anterior. ###
def guardar_contactos(lista):
    wb = openpyxl.Workbook()
    ws = wb.active
    ### Escribimos la fila de encabezado antes de agregar los datos de los contactos. ###
    ws.append(["Nombre", "Apellido", "Telefono", "Correo", "Direccion", "Categoria", "Favorito"])
    for c in lista:
        ws.append([c["nombre"], c["apellido"], c["telefono"], c["correo"], c["direccion"], c["categoria"], c["favorito"]])
    wb.save(EXCEL)


### Esta función verifica si el usuario no ha iniciado sesión, retornando True si debe ser redirigido al login. ###
def login_requerido():
    return not session.get('user')


### Esta ruta funciona como la ruta "raíz". que redirige automáticamente al login al ingresar al sistema. ###
@app.route("/")
def inicio():
    return redirect(url_for('login'))


### Esta ruta muestra el formulario en GET y valida las credenciales en POST, creando la sesión si son correctas. ###
@app.route("/login-page", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        usuario = request.form.get('usuario', '').strip()
        password = request.form.get('contrasena', '').strip()
        if not usuario or not password:
            error = "Por favor, rellena todos los campos."
        elif usuario != USUARIO:
            error = "El usuario ingresado no existe."
        elif password != PASSWORD:
            error = "La contrasena es incorrecta."
        else:
            session['user'] = usuario
            return redirect(url_for('contactos'))
    return render_template("login-page.html", error=error)


### Eliminamos el usuario de la sesión y redirigimos al login si el usuario lo desea. ###
@app.route("/logout")
def logout():
    session.pop('user', None)
    return redirect(url_for('login'))


### Esta ruta carga la lista, permite ordenarla alfabéticamente y calcula estadísticas para la vista. ###
@app.route("/contactos")
def contactos():
    if login_requerido():
        return redirect(url_for('login'))
    lista = leer_contactos()
    orden = request.args.get('orden')
    ### Si se solicita orden alfabético, ordenamos la lista por nombre en minúsculas. ###
    if orden == 'az':
        lista = sorted(lista, key=lambda x: x['nombre'].lower())
    total_favoritos = sum(1 for c in lista if c['favorito'] == 'si')
    total_trabajo = sum(1 for c in lista if c['categoria'] == 'Trabajo')
    return render_template("contacts.html", contactos=lista, total_favoritos=total_favoritos, total_trabajo=total_trabajo)


### Creamos una muestra el formulario en GET y valida y guarda los datos en POST. ###
@app.route("/nuevo-contacto", methods=["GET", "POST"])
def nuevo_contacto():
    if login_requerido():
        return redirect(url_for('login'))
    error = ""
    if request.method == "POST":
        nombre = request.form.get('nombre', '').strip()
        apellido = request.form.get('apellido', '').strip()
        telefono = request.form.get('telefono', '').strip()
        correo = request.form.get('correo', '').strip()
        direccion = request.form.get('direccion', '').strip()
        categoria = request.form.get('categoria', 'Otro')
        favorito = request.form.get('favorito', 'no')
        ### Validamos que los campos obligatorios estén completos y con formato correcto antes de guardar. ###
        if not nombre or not telefono or not correo:
            error = "Los campos nombre, telefono y correo son obligatorios."
        elif '@' not in correo or '.' not in correo.split('@')[-1]:
            error = "El correo no tiene un formato valido."
        elif not telefono.isdigit() or len(telefono) != 8:
            error = "El telefono debe tener exactamente 8 digitos."
        ### Si no hay errores, agregamos el nuevo contacto a la lista y lo guardamos en el Excel. ###
        if not error:
            lista = leer_contactos()
            lista.append({
                "nombre": nombre, "apellido": apellido, "telefono": telefono,
                "correo": correo, "direccion": direccion, "categoria": categoria, "favorito": favorito
            })
            guardar_contactos(lista)
            return redirect(url_for('contactos'))
    return render_template("nuevo_contacto.html", error=error)


### Filtramos la lista por nombre o apellido según el parámetro "q" recibido por GET. ###
@app.route("/buscar-contacto")
def buscar_contacto():
    if login_requerido():
        return redirect(url_for('login'))
    q = request.args.get('q', '').strip().lower()
    resultados = []
    if q:
        for c in leer_contactos():
            if q in c['nombre'].lower() or q in c['apellido'].lower():
                resultados.append(c)
    return render_template("buscar_contacto.html", resultados=resultados, busqueda=q)


### En este bloque, cargamos sus datos en GET y actualiza el registro en POST tras validación. ###
@app.route("/editar-contacto/<int:idx>", methods=["GET", "POST"])
def editar_contacto(idx):
    if login_requerido():
        return redirect(url_for('login'))
    lista = leer_contactos()
    if idx >= len(lista):
        return redirect(url_for('contactos'))
    contacto = lista[idx]
    error = ""
    if request.method == "POST":
        nombre = request.form.get('nombre', '').strip()
        apellido = request.form.get('apellido', '').strip()
        telefono = request.form.get('telefono', '').strip()
        correo = request.form.get('correo', '').strip()
        direccion = request.form.get('direccion', '').strip()
        categoria = request.form.get('categoria', 'Otro')
        favorito = request.form.get('favorito', 'no')
        ### Validamos los campos obligatorios y el formato del correo y teléfono antes de actualizar. ###
        if not nombre or not telefono or not correo:
            error = "Los campos nombre, telefono y correo son obligatorios."
        elif '@' not in correo or '.' not in correo.split('@')[-1]:
            error = "El correo no tiene un formato valido."
        elif not telefono.isdigit() or len(telefono) != 8:
            error = "El telefono debe tener exactamente 8 digitos."
        ### Si no hay errores, actualizamos el contacto en la posición correspondiente y guardamos el Excel. ###
        if not error:
            lista[idx] = {
                "nombre": nombre, "apellido": apellido, "telefono": telefono,
                "correo": correo, "direccion": direccion, "categoria": categoria, "favorito": favorito
            }
            guardar_contactos(lista)
            return redirect(url_for('contactos'))
        ### Si hay errores, reconstruimos el diccionario del contacto con los valores ingresados para repoblar el formulario. ###
        contacto = {"nombre": nombre, "apellido": apellido, "telefono": telefono,
                    "correo": correo, "direccion": direccion, "categoria": categoria, "favorito": favorito}
    return render_template("editar_contacto.html", contacto=contacto, idx=idx, error=error)


### Este bloque funciona como una ruta para eliminar un contacto por índice: lo retira de la lista y guarda los cambios en el Excel. ###
@app.route("/eliminar-contacto/<int:idx>")
def eliminar_contacto(idx):
    if login_requerido():
        return redirect(url_for('login'))
    lista = leer_contactos()
    if idx < len(lista):
        lista.pop(idx)
        guardar_contactos(lista)
    return redirect(url_for('contactos'))


### Caargamos sus datos por índice y los pasa a la plantilla de perfil. ###
@app.route("/perfil/<int:idx>")
def perfil_contacto(idx):
    if login_requerido():
        return redirect(url_for('login'))
    lista = leer_contactos()
    if idx >= len(lista):
        return redirect(url_for('contactos'))
    return render_template("perfil.html", contacto=lista[idx], idx=idx)


### Calculamos el total de contactos, favoritos y la cantidad por categoría para mostrarlos en la vista. ###
@app.route("/reporte")
def reporte():
    if login_requerido():
        return redirect(url_for('login'))
    lista = leer_contactos()
    total = len(lista)
    favoritos = sum(1 for c in lista if c['favorito'] == 'si')
    ### Construimos el diccionario de categorías contando cuántos contactos pertenecen a cada una. ###
    categorias = {}
    for c in lista:
        cat = c['categoria']
        categorias[cat] = categorias.get(cat, 0) + 1
    return render_template("reporte.html", total=total, favoritos=favoritos, categorias=categorias)


### En este bloque, ejecutamos el servidor en modo debug para desarrollo. ###
if __name__ == "__main__":
    app.run(debug=True)